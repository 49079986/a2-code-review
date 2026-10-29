"""Read local course materials; preserve source files and export text/page evidence."""
from pathlib import Path
import json, hashlib, zipfile, re, shutil
import pymupdf as fitz
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / 'A2'
OUT = ROOT / 'data' / 'extracted'
OUT.mkdir(parents=True, exist_ok=True)
SKIP = {'node_modules','.git','venv','.venv','__pycache__','pgdata','pg_data','base','global','pg_wal','pg_xact','pg_multixact','pg_stat','pg_subtrans','pg_tblspc','pg_twophase','pg_logical','pg_notify','pg_serial','pg_snapshots','pg_commit_ts','pg_replslot','pg_dynshmem'}
def walk(path):
    for p in sorted(path.iterdir()):
        if p.is_dir():
            if p.name.lower() not in SKIP:
                yield from walk(p)
        else:
            yield p

items=[]
for p in walk(SOURCE):
    if p.name.startswith('~$'): continue
    if p.suffix.lower() not in {'.pdf','.docx','.sql','.png','.jpg','.jpeg','.txt','.log','.md','.zip'}: continue
    if p.stat().st_size > 40_000_000: continue
    rel=p.relative_to(SOURCE).as_posix()
    key=re.sub(r'[^a-zA-Z0-9._-]','_',rel)
    item={'source':rel,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
    if p.suffix.lower()=='.pdf':
        doc=fitz.open(p)
        item['pages']=len(doc)
        txt='\n\n'.join(f'=== PDF PAGE {i+1} ===\n'+page.get_text() for i,page in enumerate(doc))
        (OUT/(key+'.txt')).write_text(txt,encoding='utf-8')
        item['text_file']='data/extracted/'+key+'.txt'
        item['images_per_page']=[len(page.get_images()) for page in doc]
    elif p.suffix.lower()=='.docx':
        d=Document(p)
        txt='\n'.join(x.text for x in d.paragraphs)+'\n\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows)
        (OUT/(key+'.txt')).write_text(txt,encoding='utf-8')
        item['text_file']='data/extracted/'+key+'.txt'
        with zipfile.ZipFile(p) as z: item['embedded_images']=[n for n in z.namelist() if n.startswith('word/media/')]
    items.append(item)
(ROOT/'data'/'source-inventory.json').write_text(json.dumps(items,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps([{k:v for k,v in i.items() if k not in ('embedded_images',)} for i in items if i['source'].endswith(('.pdf','.docx','.md','.txt','.log'))],ensure_ascii=False,indent=2))
